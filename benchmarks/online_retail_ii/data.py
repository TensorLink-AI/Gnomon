"""Stream transactions with explicit ownership, selection and visibility rules."""
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
import hashlib
import io
import json
from pathlib import Path
import re
import zipfile

SOURCE_SHA = '572e36277c2390fbfde10664750731e0a86f55e33470d91919085f0408e67bfb'
START = date(2009, 12, 7)
SELECTION_END = date(2010, 11, 30)
BOUNDARY = date(2010, 12, 1)
FIRST_ORIGIN = date(2010, 12, 5)
PHASES = {'development': range(13), 'validation': range(13, 19), 'final': range(19, 26)}
HEADERS = ('Invoice', 'StockCode', 'Description', 'Quantity', 'InvoiceDate', 'Price', 'Customer ID', 'Country')


def sha(path):
    with Path(path).open('rb') as f:
        return hashlib.file_digest(f, 'sha256').hexdigest()


def dump(path, value):
    Path(path).write_text(json.dumps(value, indent=2, sort_keys=True, allow_nan=False)+'\n')


def protocol_sha():
    return sha(Path(__file__).with_name('PROTOCOL.md'))


def origin(index):
    return FIRST_ORIGIN+timedelta(days=14*index)


def classify(row, sheet, cutoff):
    """Reject unavailable dates before using quantity/price/customer fields."""
    when = row[4]
    if not isinstance(when, datetime):
        return None, 'invalid_date'
    day = when.date()
    if day > cutoff:
        return None, 'after_phase_cutoff_not_used'
    if day < START:
        return None, 'before_complete_calendar'
    if (sheet == 'Year 2009-2010' and day >= BOUNDARY) or (sheet == 'Year 2010-2011' and day < BOUNDARY):
        return None, 'other_sheet_owns_date'
    if row[7] != 'United Kingdom':
        return None, 'other_country'
    if str(row[0]).strip().upper().startswith('C'):
        return None, 'cancellation'
    code = str(row[1]).strip().upper()
    if not re.fullmatch(r'\d{5}[A-Z]*', code):
        return None, 'non_product_stock_code'
    quantity, price = row[3], row[5]
    if type(quantity) not in (int, float) or not float(quantity).is_integer():
        return None, 'invalid_quantity'
    if quantity <= 0:
        return None, 'nonpositive_quantity'
    if type(price) not in (int, float) or not 0 < price < float('inf'):
        return None, 'nonpositive_or_invalid_price'
    return (code, day, int(quantity)), 'accepted'


def cohort(daily):
    buckets = {'sparse': [], 'intermittent': [], 'frequent': []}
    details = {}
    days = (SELECTION_END-START).days+1
    for code, values in sorted(daily.items()):
        past = {d: v for d, v in values.items() if START <= d <= SELECTION_END and v > 0}
        if not past:
            continue
        dates = sorted(past)
        weeks = len({d-timedelta(days=d.weekday()) for d in dates})
        if sum(past.values()) < 24 or weeks < 8 or (dates[-1]-dates[0]).days < 180 or (SELECTION_END-dates[-1]).days >= 56:
            continue
        fraction = len(dates)/days
        bucket = 'sparse' if fraction <= .1 else 'intermittent' if fraction <= .3 else 'frequent'
        digest = hashlib.sha256(('online-retail-ii-v1:'+code).encode()).hexdigest()
        buckets[bucket].append((digest, code))
        details[code] = {'stratum': bucket, 'active_fraction': fraction, 'active_weeks': weeks,
                         'training_units': sum(past.values()), 'selection_hash': digest}
    selected = {code: details[code] for bucket in buckets.values() for _, code in sorted(bucket)[:16]}
    return selected, {name: len(values) for name, values in buckets.items()}


def prepare(archive, output, phase='development', release=None):
    import openpyxl
    import pandas as pd
    archive, output = Path(archive), Path(output)
    if output.exists():
        raise ValueError('Use a fresh output directory; retain previous preparations')
    if phase not in PHASES or sha(archive) != SOURCE_SHA:
        raise ValueError('Wrong phase or source ZIP fingerprint')
    if phase != 'development':
        if release is None:
            raise ValueError('Validation/final requires a separately frozen release manifest')
        released = json.loads(Path(release).read_text())
        if (released.get('phase') != phase or released.get('source_sha256') != SOURCE_SHA
                or released.get('protocol_sha256') != protocol_sha()
                or not released.get('frozen_candidate_commit') or not released.get('runtime_lock_sha256')):
            raise ValueError('Release manifest does not bind this source, phase, protocol, code and runtime')
    cutoff = origin(max(PHASES[phase]))+timedelta(days=14)
    with zipfile.ZipFile(archive) as z:
        if z.namelist() != ['online_retail_II.xlsx']:
            raise ValueError('Unexpected archive member set')
        workbook_bytes = z.read('online_retail_II.xlsx')
    workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    if workbook.sheetnames != ['Year 2009-2010', 'Year 2010-2011']:
        raise ValueError('Unexpected sheet schema')
    daily = defaultdict(lambda: defaultdict(int)); counters = {}; selected = None
    try:
        for sheet in workbook:
            counts = Counter(); rows = sheet.iter_rows(values_only=True)
            if tuple(next(rows)) != HEADERS:
                raise ValueError('Unexpected column schema')
            for row in rows:
                counts['rows_scanned'] += 1
                item, reason = classify(row, sheet.title, cutoff)
                counts[reason] += 1
                if item is None:
                    continue
                code, day, quantity = item
                if selected is None or code in selected:
                    daily[code][day] += quantity
            counters[sheet.title] = dict(counts)
            if selected is None:
                selected, eligible = cohort(daily)
                if not selected:
                    raise ValueError('No training-eligible products')
                daily = {code: daily[code] for code in selected}
            print(json.dumps({'sheet_complete': sheet.title, 'selected_products': len(selected)}), flush=True)
    finally:
        workbook.close()
    output.mkdir(parents=True)
    calendar = pd.date_range(START, cutoff, freq='D')
    frame = pd.DataFrame([{'series_id': code, 'date': stamp.date().isoformat(),
                           'value': daily[code].get(stamp.date(), 0)}
                          for code in sorted(selected) for stamp in calendar])
    frame.to_csv(output/'host-panel.csv', index=False)
    manifest = {'schema': 'online-retail-ii-v1', 'phase': phase, 'source_sha256': SOURCE_SHA,
        'workbook_sha256': hashlib.sha256(workbook_bytes).hexdigest(), 'protocol_sha256': protocol_sha(),
        'panel_sha256': sha(output/'host-panel.csv'), 'selection_end': SELECTION_END.isoformat(),
        'start': START.isoformat(), 'materialized_through': cutoff.isoformat(),
        'selected_products': selected, 'eligible_per_stratum': eligible, 'series': len(selected),
        'planned_origins': [origin(i).isoformat() for i in PHASES[phase]],
        'planned_cases': len(selected)*len(PHASES[phase]), 'row_diagnostics': counters,
        'target': 'UK recorded gross positive-sale units per SKU/day', 'unit': 'units',
        'timezone': 'Europe/London', 'source_availability': 'assumed_local_end_of_day',
        'recorded_times': 'not_in_source', 'stock_availability': 'unknown',
        'repeated_rows_within_owning_sheet': 'retained', 'customer_identifiers_exported': False,
        'final_targets_materialized': phase == 'final'}
    dump(output/'manifest.json', manifest)
    return manifest


def load_panel(root):
    import pandas as pd
    root = Path(root); manifest = json.loads((root/'manifest.json').read_text())
    if sha(root/'host-panel.csv') != manifest['panel_sha256'] or manifest['protocol_sha256'] != protocol_sha():
        raise ValueError('Panel or protocol changed')
    frame = pd.read_csv(root/'host-panel.csv', dtype={'series_id': str}, parse_dates=['date'])
    return manifest, frame
