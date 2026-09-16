"""Early-cohort preparation, invoked only after predecessor terminal success."""
from collections import defaultdict, Counter
from datetime import date, timedelta
import hashlib
import io
import json
from pathlib import Path
import zipfile

from benchmarks.online_retail_ii.data import SOURCE_SHA, START, HEADERS, classify, dump, sha

FIRST_ORIGIN=date(2010,4,11)
END=date(2011,12,4)
PHASES={'full_span':range(43)}

def origin(index):return FIRST_ORIGIN+timedelta(days=14*index)
def protocol_sha():return sha(Path(__file__).with_name('PROTOCOL.md'))


def cohort(daily):
    buckets={k:[] for k in ('sparse','intermittent','frequent')};details={}
    for code,values in sorted(daily.items()):
        past={d:v for d,v in values.items() if START<=d<=FIRST_ORIGIN and v>0}
        if not past:continue
        dates=sorted(past);weeks=len({d-timedelta(days=d.weekday()) for d in dates})
        if sum(past.values())<24 or weeks<8 or (dates[-1]-dates[0]).days<84 or (FIRST_ORIGIN-dates[-1]).days>=28:continue
        fraction=len(dates)/126
        group='sparse' if fraction<=.1 else 'intermittent' if fraction<=.3 else 'frequent'
        digest=hashlib.sha256(('online-retail-ii-full-span-v2:'+code).encode()).hexdigest()
        buckets[group].append((digest,code))
        details[code]={'stratum':group,'active_fraction':fraction,'active_weeks':weeks,
            'training_units':sum(past.values()),'selection_hash':digest}
    counts={k:len(v) for k,v in buckets.items()}
    if any(n<16 for n in counts.values()):raise ValueError('Insufficient warm-up cohort: '+str(counts))
    return {code:details[code] for group in buckets.values() for _,code in sorted(group)[:16]},counts


def prepare(archive,output):
    import openpyxl
    import pandas as pd
    output=Path(output)
    if output.exists():raise ValueError('Fresh full-span panel required')
    if sha(archive)!=SOURCE_SHA:raise ValueError('Wrong source archive')
    with zipfile.ZipFile(archive) as z:
        if z.namelist()!=['online_retail_II.xlsx']:raise ValueError('Unexpected archive members')
        payload=z.read(z.namelist()[0])
    workbook=openpyxl.load_workbook(io.BytesIO(payload),read_only=True,data_only=True)
    if workbook.sheetnames!=['Year 2009-2010','Year 2010-2011']:raise ValueError('Wrong sheets')
    early=defaultdict(lambda:defaultdict(int))
    try:
        rows=workbook.worksheets[0].iter_rows(values_only=True)
        if tuple(next(rows))!=HEADERS:raise ValueError('Wrong headers')
        for row in rows:
            item,_=classify(row,workbook.worksheets[0].title,FIRST_ORIGIN)
            if item:
                code,day,value=item;early[code][day]+=value
        selected,counts=cohort(early)
        daily={code:defaultdict(int) for code in selected};diagnostics={}
        for sheet in workbook:
            rows=sheet.iter_rows(values_only=True)
            if tuple(next(rows))!=HEADERS:raise ValueError('Wrong headers')
            counter=Counter()
            for row in rows:
                item,reason=classify(row,sheet.title,END);counter['rows_scanned']+=1;counter[reason]+=1
                if item and item[0] in selected:
                    code,day,value=item;daily[code][day]+=value
            diagnostics[sheet.title]=dict(counter)
    finally:workbook.close()
    output.mkdir(parents=True)
    calendar=pd.date_range(START,END)
    frame=pd.DataFrame([{'series_id':code,'date':stamp.date().isoformat(),'value':daily[code].get(stamp.date(),0)}
                        for code in sorted(selected) for stamp in calendar])
    frame.to_csv(output/'host-panel.csv',index=False)
    manifest={'schema':'online-retail-ii-full-span-v2','phase':'full_span','source_sha256':SOURCE_SHA,
        'protocol_sha256':protocol_sha(),'panel_sha256':sha(output/'host-panel.csv'),'selection_end':str(FIRST_ORIGIN),
        'materialized_through':str(END),'selected_products':selected,'eligible_per_stratum':counts,
        'series':48,'planned_cases':2064,'planned_origins':[str(origin(i)) for i in range(43)],
        'row_diagnostics':diagnostics,'target':'UK recorded gross positive-sale units per SKU/day',
        'unit':'units','timezone':'Europe/London','source_availability':'assumed_local_end_of_day',
        'recorded_times':'not_in_source','stock_availability':'unknown','customer_identifiers_exported':False,
        'final_targets_materialized':True,'scope':'separate_full_span_exploratory_replay_not_untouched_final'}
    dump(output/'manifest.json',manifest);return manifest


def load_panel(root):
    import pandas as pd
    root=Path(root);manifest=json.loads((root/'manifest.json').read_text())
    if sha(root/'host-panel.csv')!=manifest['panel_sha256'] or protocol_sha()!=manifest['protocol_sha256']:raise ValueError('Panel or protocol drift')
    return manifest,pd.read_csv(root/'host-panel.csv',dtype={'series_id':str},parse_dates=['date'])
